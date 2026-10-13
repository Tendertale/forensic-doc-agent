"""演示案件一键演练：把「用户实际使用的数据库」里的演示案件跑完整条链路。

用途：`python -m backend.scripts.serve` 起服务后，浏览器里立刻能看到一个**已经跑完**
的完整案件（字段、损伤、文稿 Word/PDF、校核清单、按指定位置插入的部位照片），
而不是只有材料、需要自己一步步点。

与 `backend/scripts/smoke.py` 的区别：
* `smoke.py` 用一次性数据库与存储根做**端到端断言**（可随时重跑、不污染用户数据）；
* 本脚本写**用户自己的** `data/cases.db` + `storage/`，只做「把演示案件准备好」，
  不做断言式验收（验收请用 smoke.py 与 pytest）。

安全边界：照片只作为文件路径处理，**绝不读取照片内容**（N-11）；不生成任何鉴定意见（N-1）。
"""

from __future__ import annotations

import argparse
import os
from pathlib import Path

from backend.app.constants import (
    PHOTO_PLACEHOLDER_TEXT,
    PHOTO_POSITION_DEFINITIONS,
    PHOTO_POSITION_KEYS,
    PHOTO_SECTION_LABELS,
    SEED_CASE_COUNT,
)
from backend.app.database import Database
from backend.app.storage import DEFAULT_STORAGE_ROOT, CaseStorage


def _inline_shape_count(docx_path: str | Path) -> int:
    from docx import Document

    return len(Document(str(docx_path)).inline_shapes)


def _document_text(docx_path: str | Path) -> str:
    """文稿/模板全文（段落 + 表格单元格）。

    只读文档文本：不打开、不解码任何图片，照片只当文件路径处理（N-11）。
    """

    try:
        from backend.app.docgen import extract_document_text

        return extract_document_text(docx_path) or ""
    except Exception:  # noqa: BLE001 - docgen 不可用时退回段落文本（仅用于打印统计）
        from docx import Document

        return "\n".join(paragraph.text for paragraph in Document(str(docx_path)).paragraphs)


def _photo_slot_count(docx_path: str | Path) -> int:
    """统计文稿里的照片栏位**标题**数。

    模板里每个栏位都有一个标题段落，两种写法都认：
    `照片栏位：<标签>`（`docgen.build_demo_template` 生成）与
    `<标签>（<栏目>）`（`seed._augment_template_with_photo_slots` 补齐时插入）。
    只按标题行开头匹配，避免把正文里出现的部位名误判成栏位（N-12）。
    """

    lines = [line.strip() for line in _document_text(docx_path).splitlines()]
    counted = 0
    for definition in PHOTO_POSITION_DEFINITIONS:
        section = PHOTO_SECTION_LABELS.get(definition.section, definition.section)
        prefixes = ("照片栏位：" + definition.label,
                    "{0}（{1}）".format(definition.label, section))
        if any(line.startswith(prefix) for line in lines for prefix in prefixes):
            counted += 1
    return counted


def _registered_photo_slots(docx_path: str | Path) -> int:
    """统计**模板原件**里的 `{{ photo_<key> }}` 锚点数（模板，不是渲染产物）。"""

    text = _document_text(docx_path)
    anchors = {"{{ photo_" + key + " }}" for key in PHOTO_POSITION_KEYS}
    return sum(1 for anchor in anchors if anchor in text)


def _latest_template(templates_dir: str | Path) -> Path | None:
    """挑 `templates/` 里照片栏位最完整、其次最新的模板（只读模板文本，N-11）。"""

    directory = Path(templates_dir)
    if not directory.is_dir():
        return None
    best: tuple | None = None
    for path in sorted(directory.glob("*.docx")):
        if not path.is_file():
            continue
        candidate = (_registered_photo_slots(path), path.stat().st_mtime, str(path))
        if best is None or candidate[:2] > best[:2]:
            best = candidate
    return Path(best[2]) if best else None


def refresh_case_template(database, storage, case_id: int) -> str:
    """模板刷新步骤：登记的模板副本缺照片栏位时，补登一份完整模板的**新副本**。

    背景：`workflow._template_for` 优先用本案登记的模板材料原件，而演示案件早期登记的
    副本只有 14 个 `{{ photo_<key> }}` 锚点（位置目录 18 项），于是文稿长期少 4 个栏位。

    做法：从 `storage/templates/` 取栏位最完整的模板，**复制成一份新的模板材料**登记到
    本案（绝不覆盖、绝不改写任何原件，N-8），`_template_for` 按“栏位最完整”优先选中它。
    已经是完整副本时直接跳过（幂等；也避免重复登记同一 sha256 被拒）。
    """

    import base64

    template_materials = [item for item in database.list_materials(case_id)
                          if str(item.get("kind")) == "template"]
    if not template_materials:
        return "该案件没有登记模板材料，渲染将回退到 templates/ 目录（无需刷新）"

    best_registered = None
    for item in template_materials:
        relative = str(item.get("original_path") or "")
        path = storage.absolute(relative) if relative else None
        if path is None or not Path(path).is_file():
            continue
        count = _registered_photo_slots(path)
        if best_registered is None or count > best_registered[0]:
            best_registered = (count, item, path)
    if best_registered is None:
        return "登记的模板材料文件都不存在，渲染将回退到 templates/ 目录（无需刷新）"

    current_count, current_item, current_path = best_registered
    source = _latest_template(storage.templates_dir())
    if source is None:
        return "templates/ 目录下没有可用模板，跳过刷新（登记副本仍为 {0} 个栏位）".format(
            current_count)
    source_count = _registered_photo_slots(source)

    # 触发条件（任一）：① 登记副本照片栏位不齐；② 共享模板正文与登记副本不一致
    #（例如共享模板修正了正文结构）。两者都只“补登新副本”，绝不覆盖原件（N-8）。
    same_text = _document_text(current_path) == _document_text(source)
    if current_count >= len(PHOTO_POSITION_KEYS) and same_text:
        return "登记的模板副本已含 {0}/{1} 个照片栏位且与模板目录一致，无需刷新".format(
            current_count, len(PHOTO_POSITION_KEYS))
    if source_count < current_count:
        return ("templates/ 目录里的模板栏位数（{0}）少于登记副本（{1}），跳过刷新").format(
            source_count, current_count)

    blob = Path(source).read_bytes()
    material = database.add_material(case_id, {
        "kind": "template",
        "filename": Path(source).name,
        "media_type": "docx",
        "page_count": 0,
        "content_base64": base64.b64encode(blob).decode("ascii"),
        "clarity_note": "模板栏位刷新：复制自 {0}（原登记副本仅 {1} 个照片栏位；原件未被覆盖，N-8）".format(
            Path(source).name, current_count),
    }, 0, storage, actor="系统")
    return ("已补登模板材料 #{0}（复制 {1}：{2} → {3} 个照片栏位；"
            "原副本 #{4} 保持只读，N-8）").format(
        material.get("id"), Path(source).name, current_count, source_count,
        current_item.get("id"))


def refresh_shared_demo_template(storage) -> str:
    """重建 `templates/` 里**程序化生成的演示模板**（不是用户的模板原件）。

    只在目标文件确实是我们自己生成的演示模板时才会重写：判据是正文含演示模板特有的
    中文栏目字样。真实用户上传的模板原件一律跳过（N-8：模板原件不可覆盖）。
    """

    from backend.app.docgen import build_demo_template, extract_document_text
    from backend.app.constants import DEFAULT_TEMPLATE_VERSION

    target = storage.templates_dir() / f"{DEFAULT_TEMPLATE_VERSION}.docx"
    storage.templates_dir().mkdir(parents=True, exist_ok=True)

    if target.is_file():
        try:
            existing_text = extract_document_text(target) or ""
        except Exception:  # noqa: BLE001 - 读不出来就当作不是我们的模板
            existing_text = ""
        looks_like_demo = "尸表检验" in existing_text and "部位照片" in existing_text
        if not looks_like_demo:
            return "templates/{0} 看起来不是程序化演示模板，跳过重建（N-8）".format(target.name)
        probe = target.with_suffix(".probe.docx")
        build_demo_template(probe)
        fresh_text = extract_document_text(probe) or ""
        try:
            probe.unlink()
        except OSError:
            pass
        if fresh_text == existing_text:
            return "演示模板已是最新（{0}）".format(target.name)

    build_demo_template(target)
    anchors = _registered_photo_slots(target)
    return "已重建演示模板 {0}（照片栏位 {1} 个）".format(target.name, anchors)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="把演示案件跑完并写入用户数据库（离线、不联网、不读照片内容）")
    parser.add_argument("--db", default=os.environ.get("FORENSIC_DB_PATH", "data/cases.db"),
                        help="SQLite 路径（默认 FORENSIC_DB_PATH 或 data/cases.db）")
    parser.add_argument("--storage", default=os.environ.get("FORENSIC_STORAGE_ROOT", DEFAULT_STORAGE_ROOT),
                        help="存储根（默认 FORENSIC_STORAGE_ROOT 或 storage）")
    parser.add_argument("--case-index", type=int, default=1, help="处理第几个演示案件（从 1 开始）")
    parser.add_argument("--no-pdf", action="store_true", help="跳过 PDF 导出（不依赖 LibreOffice）")
    parser.add_argument("--force", action="store_true", help="即使已有文稿也重新生成一版")
    args = parser.parse_args(argv)

    from backend.app.workflow import (
        StageUnavailable,
        run_check,
        run_extraction,
        run_normalization,
        run_ocr,
        run_preprocess,
        run_render,
    )

    database = Database(str(args.db))
    database.initialize()
    storage = CaseStorage(str(args.storage))
    storage.ensure()
    # seed 用同一个存储根写材料与照片，避免「库在 A 根、文件在 B 根」。
    database.storage = storage

    # 演示模板是程序化生成的（TODO-DL-6），正文结构升级后需要重建；
    # 真实用户上传的模板原件不会被覆盖（N-8）。
    try:
        print("[template] {0}".format(refresh_shared_demo_template(storage)))
    except Exception as error:  # noqa: BLE001 - 重建失败不阻断演示链路
        print("[template][skip] 演示模板重建失败：{0}".format(error))

    if not database.list_cases(page=1, page_size=1).get("total"):
        created = database.seed(SEED_CASE_COUNT)
        print("[seed] 已建立演示案件：{0}".format(created))
    else:
        print("[seed] 数据库已有案件，沿用现有数据")

    listing = database.list_cases(page=1, page_size=100)
    items = listing.get("items") or []
    if not items:
        print("[错误] 没有任何案件可用")
        return 2
    index = max(1, min(int(args.case_index), len(items))) - 1
    case_id = int(items[index]["id"])
    case_data = database.get_case(case_id) or {}
    case_no = str(case_data.get("case_no") or "")
    print("[case] case_id={0} case_no={1}".format(case_id, case_no))

    documents = database.list_documents(case_id) or []
    if documents and not args.force:
        print("[skip] 该案件已有 {0} 版文稿；如需重跑请加 --force".format(len(documents)))
        exit_code = 0
    else:
        # 模板刷新：登记副本缺照片栏位时补登一份完整模板的新副本（绝不覆盖原件，N-8）
        try:
            print("[template] {0}".format(refresh_case_template(database, storage, case_id)))
        except Exception as error:  # noqa: BLE001 - 刷新失败不阻断演示链路
            print("[template][skip] 模板刷新失败：{0}".format(error))

        for material in case_data.get("materials") or []:
            if material.get("kind") in ("commission", "record_book"):
                try:
                    result = run_preprocess(database, storage, case_id, int(material["id"]),
                                            ["autocontrast"], 1, 0, "系统")
                    print("[preprocess] {0}：{1}".format(material.get("filename"), result.detail))
                except StageUnavailable as error:
                    print("[preprocess][skip] {0}".format(error))

        for label, runner in (
            ("ocr", lambda: run_ocr(database, storage, case_id, engine="sidecar", force=True,
                                    expected_version=0, actor="系统")),
            ("extract", lambda: run_extraction(database, storage, case_id, replace_existing=True,
                                               expected_version=0, actor="系统")),
            ("normalize", lambda: run_normalization(database, storage, case_id, target="all",
                                                    confirm_safe=False, expected_version=0, actor="系统")),
        ):
            try:
                result = runner()
                print("[{0}] {1}".format(label, result.detail))
                if label == "ocr":
                    skipped = (result.payload or {}).get("skipped_photo_ids") or []
                    print("       照片跳过（N-11，绝不识别照片内容）：{0}".format(len(skipped)))
            except StageUnavailable as error:
                print("[{0}][不可用] {1}".format(label, error))

        formats = ("docx",) if args.no_pdf else ("docx", "pdf")
        try:
            rendered = run_render(database, storage, case_id, doc_type="report",
                                  doc_status="internal_draft", formats=formats,
                                  expected_version=0, actor="系统")
            print("[render] {0}".format(rendered.detail))
            docx_absolute = (rendered.payload or {}).get("docx_absolute") or ""
            pdf_absolute = (rendered.payload or {}).get("pdf_absolute") or ""
            if docx_absolute and Path(docx_absolute).is_file():
                print("        DOCX：{0}".format(docx_absolute))
                print("        内联图片数（应为已指派且文件存在的照片数）：{0}".format(
                    _inline_shape_count(docx_absolute)))
                text = _document_text(docx_absolute)
                print("        照片栏位标题数（应为 {0}）：{1}".format(
                    len(PHOTO_POSITION_KEYS), _photo_slot_count(docx_absolute)))
                print("        「{0}」出现次数（未指派照片的栏位）：{1}".format(
                    PHOTO_PLACEHOLDER_TEXT, text.count(PHOTO_PLACEHOLDER_TEXT)))
                residual = sum(text.count(pattern) for pattern in ("{{", "{%"))
                print("        模板占位符残留（{{{{ / {{% ，应为 0）：{0}".format(residual))
                if not rendered.payload.get("unresolved"):
                    print("        未解析占位符：0")
            if pdf_absolute and Path(pdf_absolute).is_file():
                print("        PDF ：{0}".format(pdf_absolute))
        except StageUnavailable as error:
            print("[render][不可用] {0}".format(error))

        try:
            checked = run_check(database, storage, case_id, include_layout=True,
                                expected_version=0, actor="系统")
            print("[check] {0}".format(checked.detail))
        except StageUnavailable as error:
            print("[check][不可用] {0}".format(error))
        exit_code = 0

    checklist = database.list_checklist(case_id) or []
    blocking = [item for item in checklist
                if item.get("severity") == "blocking" and item.get("status") == "open"]
    fields = database.list_fields(case_id) or []
    injuries = database.list_injuries(case_id) or []
    documents = database.list_documents(case_id) or []
    print("[summary] 字段 {0}、损伤 {1}、文稿 {2}、校核 {3}（阻断 {4}）".format(
        len(fields), len(injuries), len(documents), len(checklist), len(blocking)))
    print("[next] 启动服务：python -m backend.scripts.serve          →  浏览器打开 http://127.0.0.1:9765/")
    print("       （默认端口 9765；需要别的端口用 --port N）")
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
