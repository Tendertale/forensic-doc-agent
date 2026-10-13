"""端到端离线演练：建案 → 材料 → 识别 → 提取 → 整理 → 生成 → 校核。

用法::

    python -m backend.scripts.smoke

演练内容（全部离线、无网络请求）：

1. 建立干净的演练数据库与存储目录，写入**程序化演示案件**（`seed.seed_demo_case`）；
2. 依次执行 `workflow` 的七个阶段：预处理 → 识别 → 提取 → 整理 → 生成 → 校核 → 交付；
3. 用 `python-docx` 打开真实生成的 `.docx` 并断言（R-56 / N-5 / N-11）：
   * `len(document.inline_shapes)` == 已指派位置且文件存在的照片数；
   * 正文无 `{{ }}` / `{% %}` 残留；
   * 空白（`not_recorded`）字段为“未记录”，正文与字段都不含“未见异常”；
   * 部位照片材料 `ocr_blocks` 恒为 0；未指派位置的照片不出现在任何栏位；
   * **D7**：文稿照片占位文字“（未提供照片）”的次数 == `GET /photo-slots`
     的 `missing_count`，栏位数 == `PHOTO_POSITION_DEFINITIONS` 的 18 项；
4. **D9**：第 [10] 段用**真实 HTTP + 真实最新版本号**（不再一律传 `expected_version=0`）
   重走一遍写链路——建案 → 材料（请求体与查询参数两种版本号位置）→ 部位照片 →
   预处理 → 识别 → 提取 → 整理 → 生成 → 校核 → 豁免阻断项 → 审核 → 交付，
   逐步断言状态码，使 D1（`extract` 恒 409）/ D2（`materials` 恒 422）/ D5（审核绕过阻断门）
   这类「只在真实 HTTP 路径上才暴露」的装配缺陷今后无法再被 `expected_version=0` 掩盖。
"""

from __future__ import annotations

import os
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

DB_PATH = PROJECT_ROOT / "data" / "smoke.db"
STORAGE_ROOT = PROJECT_ROOT / "storage" / "smoke"

_FAILURES: List[str] = []
_WAITING: List[str] = []


def _line(title: str, ok: bool, detail: str = "") -> None:
    mark = "OK  " if ok else "FAIL"
    print("[{0}] {1}{2}".format(mark, title, ("：" + detail) if detail else ""))
    if not ok:
        _FAILURES.append(title)


def _check(title: str, condition: bool, detail: str = "") -> bool:
    _line(title, bool(condition), detail if condition else (detail or "断言不成立"))
    return bool(condition)


def _reset_environment() -> None:
    """清空上一次演练的数据库与存储目录（仅限本脚本专用的 smoke 路径）。"""

    import shutil

    for target in (DB_PATH, Path(str(DB_PATH) + "-journal"), Path(str(DB_PATH) + "-wal")):
        try:
            if target.is_file():
                target.unlink()
        except OSError:
            pass
    marker = str(STORAGE_ROOT.resolve())
    if marker.startswith(str(PROJECT_ROOT.resolve())) and STORAGE_ROOT.name == "smoke":
        shutil.rmtree(STORAGE_ROOT, ignore_errors=True)


def _pdf_pages(path: Path) -> Optional[int]:
    """PDF 页数（PyMuPDF / pypdfium2 任一可用）；不可用时返回 None。"""

    try:
        import fitz  # type: ignore

        with fitz.open(str(path)) as document:
            return int(document.page_count)
    except Exception:  # noqa: BLE001 - 退回 pypdfium2
        pass
    try:
        import pypdfium2 as pdfium  # type: ignore

        pdf = pdfium.PdfDocument(str(path))
        return len(pdf)
    except Exception:  # noqa: BLE001
        return None


def _docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _case_version(client, case_id: int) -> int:
    """当前案件版本号：所有写操作 `expected_version` 的**真实来源**（D9 新增）。"""

    response = client.get("/api/cases/{0}".format(case_id))
    assert response.status_code == 200, response.text
    return int(response.json()["version"])


def _http_step(client, label: str, method: str, url: str, *, json_body=None, params=None):
    """执行一次真实 HTTP 写操作并断言 200，打印状态码与 detail（D9）。"""

    response = client.request(method, url, json=json_body, params=params)
    detail = ""
    try:
        payload = response.json()
    except Exception:  # noqa: BLE001 - 非 JSON 响应不阻断演练
        payload = {}
    if isinstance(payload, dict):
        detail = str(payload.get("detail") or "")
    _check("HTTP {0} {1}".format(method.upper(), label), response.status_code == 200,
           "状态码 {0}{1}".format(response.status_code, ("；" + detail[:80]) if detail else ""))
    return response


def _waive_open_blocking_db(database, case_id: int) -> List[str]:
    """直接在数据层豁免本案件的 open blocking 项（D5 之后 `approve` 要求没有 open 阻断项）。"""

    waived: List[str] = []
    for item in database.open_blocking_items(case_id):
        version = int((database.get_case(case_id) or {}).get("version") or 1)
        database.update_checklist_item(int(item["id"]), {
            "status": "waived", "resolved_by": "演练",
            "note": "演练豁免：为验证交付链路（D5 后 approve 要求无 open blocking）",
        }, version, "演练")
        waived.append(str(item.get("code")))
    return waived


def _waive_open_blocking_http(client, case_id: int) -> List[str]:
    """经 HTTP 豁免本案件的 open blocking 项（每一步都带真实版本号）。"""

    waived: List[str] = []
    for item in client.get("/api/cases/{0}/checklist".format(case_id)).json():
        if item.get("severity") != "blocking" or item.get("status") != "open":
            continue
        response = client.patch("/api/checklist/{0}".format(item["id"]), json={
            "expected_version": _case_version(client, case_id),
            "status": "waived",
            "note": "演练豁免：HTTP 真实版本号链路校验（D5 后 approve 要求无 open blocking）",
        })
        _check("HTTP 豁免阻断项 {0}".format(item.get("code")), response.status_code == 200,
               "状态码 {0}".format(response.status_code))
        waived.append(str(item.get("code")))
    return waived


def main() -> int:
    print("=" * 78)
    print("法医尸表检验文书整理 Agent —— 端到端离线演练（不联网、不读照片内容）")
    print("=" * 78)

    _reset_environment()

    from fastapi.testclient import TestClient

    from backend.app.constants import PHOTO_PLACEHOLDER_TEXT, PHOTO_POSITION_DEFINITIONS
    from backend.app.database import utc_now
    from backend.app.main import create_app
    from backend.app.seed import (
        demo_commission_blocks,
        demo_scan_bytes,
        placeholder_photo_bytes,
        seed_demo_case,
    )
    from backend.app.workflow import (
        StageUnavailable,
        photo_context,
        run_check,
        run_delivery,
        run_extraction,
        run_normalization,
        run_ocr,
        run_preprocess,
        run_render,
    )

    app = create_app(str(DB_PATH), str(STORAGE_ROOT), seed_data=False)
    database = app.state.db
    storage = app.state.storage

    # ------------------------------------------------------------------
    print("\n[1] 应用与接口自检")
    # ------------------------------------------------------------------
    client = TestClient(app)
    health = client.get("/api/health")
    _check("GET /api/health", health.status_code == 200 and health.json()["status"] == "ok",
           str(health.json()))
    config = client.get("/api/config")
    _check("GET /api/config", config.status_code == 200,
           "photo_positions={0} 个、阶段 {1} 个、demo_data={2}".format(
               len(config.json().get("photo_positions") or []),
               len(config.json().get("stage_order") or []),
               config.json().get("demo_data")))
    root = client.get("/")
    _check("GET / 工作台/回退页", root.status_code == 200,
           root.headers.get("content-type", ""))

    # ------------------------------------------------------------------
    print("\n[2] 建案与材料（演示数据，TODO-DL-6 / TODO-DL-7）")
    # ------------------------------------------------------------------
    case_id = seed_demo_case(database, storage, index=1)
    case_data = database.get_case(case_id) or {}
    case_no = str(case_data.get("case_no"))
    kinds = [item["kind"] for item in case_data.get("materials") or []]
    print("     案件编号：{0}（case_id={1}）".format(case_no, case_id))
    print("     案件目录：{0}".format(storage.case_dir(case_no)))
    _check("材料登记", {"commission", "record_book", "template", "photo"} <= set(kinds),
           "kind=" + "、".join(sorted(set(kinds))))
    photo_materials = [item for item in case_data.get("materials") or []
                       if item["kind"] == "photo"]
    _check("部位照片数量 ≥ 4", len(photo_materials) >= 4, "{0} 张".format(len(photo_materials)))
    events = database.list_events(case_id)
    _check("intake 事件", any(item["event_type"] == "intake" for item in events),
           "、".join(sorted({str(item["event_type"]) for item in events}))[:90])

    # ------------------------------------------------------------------
    print("\n[3] 预处理 → 识别（N-11：照片一律跳过）")
    # ------------------------------------------------------------------
    for material in case_data.get("materials") or []:
        if material["kind"] in ("commission", "record_book"):
            try:
                result = run_preprocess(database, storage, case_id, int(material["id"]),
                                        ["autocontrast"], 1, 0, "系统")
                print("     {0}：{1}".format(material["filename"], result.detail))
            except StageUnavailable as error:
                _WAITING.append("workflow.run_preprocess → {0}".format(error))

    try:
        ocr_result = run_ocr(database, storage, case_id, engine="sidecar", force=True,
                             expected_version=0, actor="系统")
        print("     {0}".format(ocr_result.detail))
        _check("OCR 完成（sidecar）", ocr_result.payload.get("blocks", 0) > 0,
               "写入 {0} 个文本块，跳过照片 {1} 张".format(
                   ocr_result.payload.get("blocks"),
                   len(ocr_result.payload.get("skipped_photo_ids") or [])))
        _check("R-55 照片列入 skipped_photo_ids",
               len(ocr_result.payload.get("skipped_photo_ids") or []) >= 4,
               str(ocr_result.payload.get("skipped_photo_ids")))
    except StageUnavailable as error:
        _WAITING.append("workflow.run_ocr → {0}".format(error))

    photo_block_total = 0
    for material in photo_materials:
        photo_block_total += int(database.list_ocr_blocks(
            case_id, int(material["id"]), page=1, page_size=200)["total"])
    _check("N-11 照片材料 ocr_blocks 恒为 0", photo_block_total == 0,
           "照片 {0} 张，共 {1} 个识别块".format(len(photo_materials), photo_block_total))

    # ------------------------------------------------------------------
    print("\n[4] 提取 → 整理")
    # ------------------------------------------------------------------
    try:
        extract_result = run_extraction(database, storage, case_id, replace_existing=True,
                                       expected_version=0, actor="系统")
        print("     {0}".format(extract_result.detail))
        _check("字段提取", extract_result.payload.get("fields", 0) > 0,
               "字段 {0} 项、损伤 {1} 条、冲突 {2} 处".format(
                   extract_result.payload.get("fields"),
                   extract_result.payload.get("injuries"),
                   len(extract_result.payload.get("conflicts") or [])))
        _check("R-16 损伤逐条", extract_result.payload.get("injuries", 0) >= 2,
               "{0} 条".format(extract_result.payload.get("injuries")))
    except StageUnavailable as error:
        _WAITING.append("workflow.run_extraction → {0}".format(error))

    try:
        normalize_result = run_normalization(database, storage, case_id, target="all",
                                             expected_version=0, actor="系统")
        print("     {0}".format(normalize_result.detail))
        _check("术语整理", normalize_result.payload.get("fields", 0) > 0,
               "字段 {0} 项、损伤 {1} 条、待确认 {2} 项".format(
                   normalize_result.payload.get("fields"),
                   normalize_result.payload.get("injuries"),
                   normalize_result.payload.get("requires_confirmation")))
    except StageUnavailable as error:
        _WAITING.append("workflow.run_normalization → {0}".format(error))

    # ------------------------------------------------------------------
    print("\n[5] 照片栏位（R-54 / R-56 / R-57）")
    # ------------------------------------------------------------------
    slots = photo_context(database, storage, case_id)
    assigned = [photo for slot in slots for photo in slot.get("photos") or []]
    usable = [photo for photo in assigned if Path(str(photo.get("path"))).is_file()]
    unassigned = [item for item in photo_materials
                  if not str(item.get("position_key") or "")]
    print("     已指派栏位：" + "、".join(
        "{0}×{1}".format(slot["label"], len(slot["photos"]))
        for slot in slots if slot["photos"]))
    print("     未指派照片：{0} 张（不进入文稿，触发 PHOTO_POSITION_MISSING）".format(
        len(unassigned)))
    _check("已指派照片 ≥ 3 且分属不同栏位",
           len(assigned) >= 3 and len({slot["position_key"] for slot in slots
                                       if slot["photos"]}) >= 3,
           "{0} 张 / {1} 个栏位".format(len(assigned),
                                       len({slot["position_key"] for slot in slots
                                            if slot["photos"]})))
    _check("R-57 未指派照片不出现在任何栏位",
           all(photo["material_id"] not in {item["id"] for item in unassigned}
               for photo in assigned),
           "未指派 {0} 张".format(len(unassigned)))

    # ------------------------------------------------------------------
    print("\n[6] 生成文稿（render）")
    # ------------------------------------------------------------------
    try:
        render_result = run_render(database, storage, case_id, doc_type="report",
                                   doc_status="internal_draft", formats=("docx", "pdf"),
                                   expected_version=0, actor="系统")
        payload = render_result.payload
        print("     {0}".format(render_result.detail))
        print("     DOCX：{0}".format(payload.get("docx_absolute")))
        print("     PDF ：{0}".format(payload.get("pdf_absolute") or "（未生成）"))
        document_id = int(payload.get("document_id") or 0)
    except StageUnavailable as error:
        _WAITING.append("workflow.run_render → {0}".format(error))
        document_id = 0
        payload = {}

    if document_id:
        document = database.get_document(document_id) or {}
        docx_path = storage.absolute(str(document.get("docx_path") or ""))
        _check("DOCX 文件真实存在", docx_path.is_file(), str(docx_path))

        from docx import Document

        doc = Document(str(docx_path))
        shapes = len(doc.inline_shapes)
        print('     断言：len(document.inline_shapes) == 已指派且文件存在的照片数')
        _check("R-56 inline_shapes == 可用照片数", shapes == len(usable),
               "inline_shapes={0}，已指派可用照片={1}".format(shapes, len(usable)))

        text = _docx_text(docx_path)
        leftovers = [token for token in ("{{", "}}", "{%", "%}") if token in text]
        _check("正文无占位符残留", not leftovers,
               "残留：" + "、".join(leftovers) if leftovers else "未发现 {{{{ }}}} / {{% %}}")
        _check("正文不含“未见异常”", "未见异常" not in text,
               "命中" if "未见异常" in text else "未命中")

        # D7 / N-12：文稿照片占位文字口径必须与 GET /photo-slots 闭合
        slots_api = client.get("/api/cases/{0}/photo-slots".format(case_id)).json()
        placeholders = text.count(PHOTO_PLACEHOLDER_TEXT)
        _check("D7/N-12 文稿照片占位文字数 == photo-slots 空栏位数",
               placeholders == int(slots_api.get("missing_count") or 0),
               "占位 {0} 处 / missing_count={1} / 栏位 {2} 个".format(
                   placeholders, slots_api.get("missing_count"),
                   len(slots_api.get("slots") or [])))
        _check("D7 照片栏位目录与文稿字段目录一致",
               len(slots_api.get("slots") or []) == len(PHOTO_POSITION_DEFINITIONS),
               "{0} 个栏位 / 目录 {1} 项".format(
                   len(slots_api.get("slots") or []), len(PHOTO_POSITION_DEFINITIONS)))

        pdf_path = storage.absolute(str(document.get("pdf_path") or ""))
        if pdf_path.is_file():
            pages = _pdf_pages(pdf_path)
            _check("PDF 与 DOCX 同版本且页数 ≥ 1", (pages or 0) >= 1,
                   "PDF 页数={0}".format(pages))
        else:
            _WAITING.append("docgen.export_pdf（LibreOffice 未产出 PDF）")

    # ------------------------------------------------------------------
    print("\n[7] 字段状态复查（N-5 / INV-3）")
    # ------------------------------------------------------------------
    fields = database.list_fields(case_id)
    not_recorded = [item for item in fields if item["status"] == "not_recorded"]
    bad = [item for item in not_recorded
           if item["normalized_text"] != "未记录" or "未见异常" in item["normalized_text"]]
    _check("not_recorded 一律“未记录”且不含“未见异常”", not bad,
           "空白字段 {0} 个：{1}".format(len(not_recorded),
                                        "、".join(item["field_key"] for item in not_recorded))
           + ("；异常：" + "、".join(item["field_key"] for item in bad) if bad else ""))
    unclear = [item for item in fields if item["status"] == "unclear"]
    print("     待辨认/冲突字段：{0} 个".format(len(unclear)))
    conflicts = [item for item in fields if item.get("conflict")]
    print("     跨材料冲突字段：{0} 个".format(len(conflicts)))

    # ------------------------------------------------------------------
    print("\n[8] 校核（check → 清单摘要）")
    # ------------------------------------------------------------------
    try:
        check_result = run_check(database, storage, case_id,
                                 document_id=document_id or None, include_layout=True,
                                 expected_version=0, actor="系统")
        checklist = database.list_checklist(case_id)
        severity_counts: Dict[str, int] = {}
        for item in checklist:
            severity_counts[item["severity"]] = severity_counts.get(item["severity"], 0) + 1
        print("     校核清单：{0}".format(
            "、".join("{0} {1}".format(key, value)
                      for key, value in sorted(severity_counts.items())) or "空"))
        for item in checklist[:8]:
            print("       - [{0}/{1}] {2}：{3}".format(
                item["severity"], item["status"], item["code"], item["title"]))
        _check("校核清单已写入", len(checklist) > 0, "{0} 条".format(len(checklist)))
        _check("run_check 有输出", check_result.ok, check_result.detail)
    except StageUnavailable as error:
        _WAITING.append("workflow.run_check → {0}".format(error))

    # ------------------------------------------------------------------
    print("\n[9] 交付（仅已审核文稿，INV-8；D5：open blocking 项未处理时不得审核）")
    # ------------------------------------------------------------------
    if document_id:
        # D5 / R-28：存在 open blocking 项时 `approve` 必须被拒绝（409）
        open_blocking_before = database.open_blocking_items(case_id)
        if open_blocking_before:
            refused = client.post("/api/documents/{0}/approve".format(document_id), json={
                "expected_version": _case_version(client, case_id)})
            refused_detail = str(refused.json().get("detail") or "")
            _check("D5/R-28 存在 open blocking 项时拒绝审核（409）",
                   refused.status_code == 409 and "存在未解决的阻断项" in refused_detail,
                   "状态码 {0}；detail={1}".format(refused.status_code, refused_detail[:70]))
            print("     open blocking 项：{0} 条".format(len(open_blocking_before)))

        waived = _waive_open_blocking_db(database, case_id)
        if waived:
            print("     为走通交付链路已豁免阻断项：{0}".format("、".join(waived[:8])))
        case_version = int((database.get_case(case_id) or {}).get("version") or 1)
        approved = database.approve_document(document_id, case_version, "演练审核", "鉴定人")
        _check("D5 阻断项豁免后可审核（200 → approved）",
               approved.get("doc_status") == "approved", str(approved.get("doc_status")))
        print("     文稿状态：{0}".format(approved.get("doc_status")))
        delivery = run_delivery(database, storage, case_id, document_id=document_id,
                                formats=("docx", "pdf"), expected_version=0, actor="复核人")
        _check("交付保存（已审核文稿）", delivery.ok, delivery.detail)
        print("     交付 DOCX：{0}".format(
            storage.absolute(str((database.get_document(document_id) or {}).get("docx_path") or ""))))
        print("     交付 PDF ：{0}".format(
            storage.absolute(str((database.get_document(document_id) or {}).get("pdf_path") or ""))))

    # ------------------------------------------------------------------
    print("\n[10] HTTP + 真实版本号的写链路校验（D1 / D2 / D5 / D9）")
    print("     说明：本段全部走真实 HTTP 端点，且每一步都用 GET /api/cases/{id} 读到的")
    print("     **真实最新版本号**；不再像旧版演练那样一律传 expected_version=0。")
    # ------------------------------------------------------------------
    import base64
    import json as _json

    created_case = client.post("/api/cases", json={
        "commission_unit": "某市公安局某分局",
        "commission_matter": "对演练对象进行尸表检验",
        "commission_date": "2026-10-11",
        "case_summary": "HTTP 真实版本号链路演练（程序化演示数据，不含真实案件）",
        "subject_name": "演练对象",
    })
    _check("HTTP 建案", created_case.status_code == 200,
           "状态码 {0}".format(created_case.status_code))
    http_case = created_case.json()
    http_case_id = int(http_case["id"])
    http_case_no = str(http_case["case_no"])
    print("     HTTP 演练案件：{0}（case_id={1}，version={2}）".format(
        http_case_no, http_case_id, http_case.get("version")))

    # D2：材料登记的两种版本号位置（请求体 = 前端契约路径；查询参数 = 兼容路径）
    body_material = _http_step(
        client, "材料登记（expected_version 放请求体）", "post",
        "/api/cases/{0}/materials".format(http_case_id),
        json_body={
            "expected_version": _case_version(client, http_case_id),
            "kind": "commission", "filename": "http_commission.png",
            "media_type": "png", "page_count": 1,
            "content_base64": base64.b64encode(demo_scan_bytes(
                "委托书", [("委托单位：某市公安局某分局", "printed")])).decode("ascii"),
        })
    commission_material_id = int(body_material.json()["id"]) if body_material.status_code == 200 else 0
    query_material = _http_step(
        client, "材料登记（expected_version 放查询参数）", "post",
        "/api/cases/{0}/materials".format(http_case_id),
        params={"expected_version": _case_version(client, http_case_id)},
        json_body={
            "kind": "record_book", "filename": "http_record_book.png",
            "media_type": "png", "page_count": 1,
            "content_base64": base64.b64encode(demo_scan_bytes(
                "尸表检验记录", [("姓名：演练对象", "handwritten")])).decode("ascii"),
        })
    photo_material = _http_step(
        client, "部位照片登记（纯色占位图 + position_key，N-11）", "post",
        "/api/cases/{0}/materials".format(http_case_id),
        json_body={
            "expected_version": _case_version(client, http_case_id),
            "kind": "photo", "filename": "http_photo_head.png", "media_type": "png",
            "content_base64": base64.b64encode(
                placeholder_photo_bytes("演示占位图：头部（HTTP 段）", 0)).decode("ascii"),
            "position_key": "head", "photo_caption": "纯色占位图（不含人体内容）",
            "sort_order": 10,
        })
    materials_http = client.get("/api/cases/{0}/materials".format(http_case_id)).json()
    _check("D2 材料经 HTTP 落库（commission + record_book + photo）",
           {item["kind"] for item in materials_http} >= {"commission", "record_book", "photo"},
           "kind=" + "、".join(sorted(item["kind"] for item in materials_http)))
    _check("N-11 照片只登记位置、不做识别（ocr_blocks = 0）",
           all(database.list_ocr_blocks(http_case_id, int(item["id"]))["total"] == 0
               for item in materials_http if item["kind"] == "photo"),
           "照片 {0} 张".format(len([item for item in materials_http
                                    if item["kind"] == "photo"])))

    if commission_material_id:
        _http_step(client, "预处理（生成识别用副本）", "post",
                   "/api/cases/{0}/materials/{1}/preprocess".format(
                       http_case_id, commission_material_id),
                   json_body={"expected_version": _case_version(client, http_case_id),
                              "operations": ["autocontrast"], "page_number": 1})
        # sidecar 引擎的离线识别结果：与识别用副本同目录同主名（架构 5.4）
        derived = storage.derived_dir(http_case_no, "commission")
        derived.mkdir(parents=True, exist_ok=True)
        (derived / "http_commission_p1.ocr.json").write_text(
            _json.dumps(demo_commission_blocks(http_case_no), ensure_ascii=False),
            encoding="utf-8")

    _http_step(client, "识别（sidecar，照片跳过）", "post",
               "/api/cases/{0}/ocr".format(http_case_id),
               json_body={"expected_version": _case_version(client, http_case_id),
                          "engine": "sidecar", "force": True})
    # D1：入口 claim 之后内部写必须传 0，否则这里恒 409
    extract_response = _http_step(
        client, "提取（D1 缺陷点：真实版本号）", "post",
        "/api/cases/{0}/extract".format(http_case_id),
        json_body={"expected_version": _case_version(client, http_case_id)})
    if extract_response.status_code == 200:
        payload = extract_response.json().get("payload") or {}
        _check("D1 提取结果有字段与来源", int(payload.get("fields") or 0) > 0,
               "字段 {0}、损伤 {1}、blocks {2}".format(
                   payload.get("fields"), payload.get("injuries"), payload.get("blocks")))
    _http_step(client, "整理（术语）", "post",
               "/api/cases/{0}/normalize".format(http_case_id),
               json_body={"expected_version": _case_version(client, http_case_id)})
    rendered = _http_step(
        client, "生成文稿（render，internal_draft）", "post",
        "/api/cases/{0}/documents".format(http_case_id),
        json_body={"expected_version": _case_version(client, http_case_id),
                   "doc_status": "internal_draft", "formats": ["docx"]})
    http_document_id = (int((rendered.json().get("payload") or {}).get("document_id") or 0)
                        if rendered.status_code == 200 else 0)
    _http_step(client, "校核（check）", "post",
               "/api/cases/{0}/check".format(http_case_id),
               json_body={"expected_version": _case_version(client, http_case_id),
                          "document_id": http_document_id or None})
    _waive_open_blocking_http(client, http_case_id)
    if http_document_id:
        approved = _http_step(
            client, "审核（approve，D5 门槛已处理阻断项）", "post",
            "/api/documents/{0}/approve".format(http_document_id),
            json_body={"expected_version": _case_version(client, http_case_id)})
        if approved.status_code == 200:
            _check("D5 审核后状态为 approved",
                   approved.json().get("doc_status") == "approved",
                   str(approved.json().get("doc_status")))
        _http_step(client, "交付（deliver，INV-8）", "post",
                   "/api/cases/{0}/deliver".format(http_case_id),
                   json_body={"expected_version": _case_version(client, http_case_id),
                              "formats": ["docx"]})
        delivered = [item for item in client.get(
            "/api/cases/{0}/documents".format(http_case_id)).json()
            if item.get("doc_status") == "approved"]
        _check("交付仅引用 approved 文稿（INV-8）", bool(delivered),
               "approved 文稿 {0} 份".format(len(delivered)))

    # 负向校验：乐观锁在真实 HTTP 路径上仍然生效（过期版本必须 409）
    stale = client.post("/api/cases/{0}/extract".format(http_case_id), json={
        "expected_version": max(1, _case_version(client, http_case_id) - 1)})
    _check("乐观锁 INV-9：过期版本号的写操作 → 409",
           stale.status_code == 409,
           "状态码 {0}；detail={1}".format(stale.status_code, str(stale.json().get("detail"))[:60]))

    # ------------------------------------------------------------------
    print("\n" + "=" * 78)
    if _WAITING:
        print("等待的依赖（阶段被跳过，不算失败）：")
        for item in _WAITING:
            print("  - " + item)
    if _FAILURES:
        print("演练失败：{0} 项断言未通过".format(len(_FAILURES)))
        for item in _FAILURES:
            print("  - " + item)
        print("=" * 78)
        return 1
    print("演练全部通过：{0}".format(utc_now()))
    print("=" * 78)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
